"""Regressões de upload e exportação com dados inteiramente sintéticos."""
import base64
from io import BytesIO
import json
import os
from pathlib import Path
import re
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from openpyxl.utils.cell import range_to_tuple
from streamlit.dataframe_util import convert_arrow_bytes_to_pandas_df
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
SAMPLE = (ROOT / 'examples' / 'atendimentos_sinteticos.csv').read_bytes()
REPLACEMENT = (
    'Id. Atendimento;Data;Identificação;Propriedades\n'
    'NOVO-SIM-001;03/03/2025 10:00:00;usuario-sintetico-novo;tag_motivo:troca |\n'
    'NOVO-SIM-002;04/03/2025 11:00:00;usuario-sintetico-novo;tag_motivo:troca |\n'
).encode('latin1')


class Upload(BytesIO):
    def __init__(self, data, name='entrada.csv'):
        super().__init__(data)
        self.name = name


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='bot-analytics-test-')
        self.previous_directory = Path.cwd()
        os.chdir(self.temporary.name)
        self.uploader = patch('streamlit.file_uploader', return_value=[Upload(SAMPLE)])
        self.mock_upload = self.uploader.start()
        self.app = AppTest.from_file(str(ROOT / 'app.py'))
        self.run_app()

    def tearDown(self):
        self.uploader.stop()
        os.chdir(self.previous_directory)
        self.temporary.cleanup()

    def run_app(self):
        self.app.run(timeout=30)
        self.assertFalse(self.app.exception, [e.message for e in self.app.exception])

    def process(self):
        self.app.button(key='processar').click()
        self.run_app()

    def add_analysis(self, segmented=False):
        self.app.selectbox(key='nova_tag').select('tag_motivo')
        self.run_app()
        if segmented:
            self.app.checkbox(key='nova_segmentar').check()
            self.run_app()
            self.app.selectbox(key='nova_col_segmento').select('Período')
            self.run_app()
        self.app.button(key='add_analise').click()
        self.run_app()

    def generate(self):
        self.app.button(key='gerar_excel').click()
        self.run_app()
        return self.app.session_state['excel_bytes'].getvalue()

    def download(self):
        links = [m.value for m in self.app.markdown if 'base64,' in m.value]
        self.assertEqual(len(links), 1)
        payload = re.search(r'base64,([^"\s]*)', links[0]).group(1)
        return base64.b64decode(payload)

    def reference_values(self, workbook, formula):
        sheet, (first_col, first_row, last_col, last_row) = range_to_tuple(formula)
        return [
            cell.value
            for row in workbook[sheet].iter_rows(
                min_row=first_row, max_row=last_row,
                min_col=first_col, max_col=last_col,
            )
            for cell in row
        ]

    def test_synthetic_flow_and_repeatable_download(self):
        self.process()
        self.add_analysis(segmented=True)
        excel = self.generate()
        workbook = load_workbook(BytesIO(excel))
        self.assertEqual(workbook.sheetnames, ['atendimentos', 'U.U e Rec', 'resumo'])
        self.assertEqual(workbook['atendimentos'].max_row, 9)  # 8 IDs + cabeçalho.
        metrics = list(workbook['U.U e Rec'].values)
        self.assertEqual(metrics[1], ('jan/25', 5, 2, '150.00%', '2.50'))
        self.assertEqual(metrics[2], ('fev/25', 3, 2, '50.00%', '1.50'))
        summary = list(workbook['resumo'].values)
        self.assertEqual(summary[1], ('tag_motivo', 'fev/25', 'jan/25', 'TOTAL'))
        self.assertEqual(summary[2], ('dúvida', 2, 2, 4))
        self.assertEqual(summary[3], ('pedido', 1, 3, 4))
        self.assertEqual(summary[4], ('TOTAL', 3, 5, 8))
        self.assertEqual(len(workbook['resumo']._charts), 1)
        for series, expected in zip(workbook['resumo']._charts[0].series, ([2, 1], [2, 3])):
            self.assertEqual(self.reference_values(workbook, series.cat.numRef.f), ['dúvida', 'pedido'])
            self.assertEqual(self.reference_values(workbook, series.val.numRef.f), expected)
        self.assertEqual(self.download(), excel)
        self.run_app()
        self.assertEqual(self.download(), excel)

    def test_interface_chart_excludes_aggregate_row_and_column(self):
        self.process()
        self.add_analysis(segmented=True)
        self.generate()
        chart = self.app.get('vega_lite_chart')[0].proto
        encoding = json.loads(chart.spec)['encoding']
        data = convert_arrow_bytes_to_pandas_df(chart.datasets[0].data.data)
        category = encoding['x']['field']
        value = encoding['y']['field']
        segment = encoding['color']['field']
        self.assertEqual(data.groupby(category)[value].sum().to_dict(), {'dúvida': 4, 'pedido': 4})
        self.assertEqual(set(data[segment]), {'fev/25', 'jan/25'})
        self.assertEqual(self.app.session_state['analises'][0]['tabela']['TOTAL'].tolist(), [4, 4, 8])

    def test_excel_keeps_categories_containing_total_and_correct_chart_references(self):
        for category in ('perda total', 'PERDA TOTAL'):
            with self.subTest(category=category):
                records = 'Id. Atendimento;Data;Identificação;Propriedades\n'
                for number, (date, motive) in enumerate([
                    ('05/01/2025', category), ('06/01/2025', 'outra categoria'),
                    ('05/02/2025', category), ('06/02/2025', 'outra categoria'),
                ], start=1):
                    records += f'TOTAL-SIM-{number};{date};usuario-sintetico-{number};tag_motivo:{motive} |\n'
                self.mock_upload.return_value = [Upload(records.encode('latin1'))]
                self.run_app()
                self.process()
                self.add_analysis(segmented=True)
                workbook = load_workbook(BytesIO(self.generate()))
                summary = workbook['resumo']
                categories = sorted([category, 'outra categoria'])
                self.assertEqual([summary.cell(row, 1).value for row in (3, 4, 5)], categories + ['TOTAL'])
                self.assertFalse(summary['A3'].font.bold)
                self.assertFalse(summary['A4'].font.bold)
                self.assertTrue(summary['A5'].font.bold)
                self.assertEqual(summary['A5'].fill.fgColor.rgb, '00EEEEEE')
                chart = summary._charts[0]
                self.assertEqual(len(chart.series), 2)
                for series, period in zip(chart.series, ('fev/25', 'jan/25')):
                    self.assertEqual(self.reference_values(workbook, series.tx.strRef.f), [period])
                    self.assertEqual(self.reference_values(workbook, series.cat.numRef.f), categories)
                    self.assertEqual(self.reference_values(workbook, series.val.numRef.f), [1, 1])

    def test_upload_preserves_local_file_and_can_be_processed_twice(self):
        sentinel = Path('entrada.csv')
        original = b'arquivo local preexistente'
        sentinel.write_bytes(original)
        self.process()
        self.assertEqual(sentinel.read_bytes(), original)
        self.process()
        self.assertEqual(sentinel.read_bytes(), original)
        self.assertEqual(len(self.app.session_state['df']), 8)

    def test_multiple_uploads_with_same_name_merge_and_deduplicate(self):
        lines = SAMPLE.splitlines(keepends=True)
        self.mock_upload.return_value = [
            Upload(b''.join(lines[:5])),
            Upload(lines[0] + b''.join(lines[5:])),
        ]
        self.run_app()
        self.process()
        self.assertEqual(len(self.app.session_state['df']), 8)
        self.assertEqual(
            self.app.session_state['df_resultados']['Total de atendimentos'].tolist(),
            [5, 3],
        )

    def test_same_filename_new_content_clears_and_refreshes_export(self):
        self.process()
        self.add_analysis()
        old_excel = self.generate()
        self.mock_upload.return_value = [Upload(REPLACEMENT)]
        self.run_app()
        state = self.app.session_state
        for key in ('df', 'analises', 'excel_bytes', 'download_ready'):
            self.assertNotIn(key, state)
        self.process()
        self.add_analysis()
        new_excel = self.generate()
        self.assertNotEqual(old_excel, new_excel)
        workbook = load_workbook(BytesIO(new_excel))
        rows = list(workbook['atendimentos'].values)
        self.assertEqual([row[0] for row in rows[1:]], ['NOVO-SIM-001', 'NOVO-SIM-002'])
        self.assertEqual(list(workbook['U.U e Rec'].values)[1], ('mar/25', 2, 1, '100.00%', '2.00'))
        self.assertEqual(self.download(), new_excel)
        self.mock_upload.return_value = []
        self.run_app()
        self.assertNotIn('df', self.app.session_state)
        self.assertNotIn('excel_bytes', self.app.session_state)

    def test_analysis_changes_invalidate_export(self):
        self.process()
        self.add_analysis()
        self.generate()
        self.add_analysis()
        self.assertNotIn('excel_bytes', self.app.session_state)
        self.generate()
        self.app.button(key='remover_0').click()
        self.run_app()
        self.assertNotIn('excel_bytes', self.app.session_state)


if __name__ == '__main__':
    unittest.main()
